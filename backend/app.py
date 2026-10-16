import base64
import hashlib
import io
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

import cv2
import fitz
import numpy as np
import pytesseract
from flask import Flask, jsonify, request
from flask_cors import CORS
from PIL import Image, ExifTags, UnidentifiedImageError

app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = 100 * 1024 * 1024

BASE_DIR = Path(__file__).resolve().parent
LOG_PATH = BASE_DIR / "evidence_log.jsonl"

IMAGE_EXTENSIONS = {
    ".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"
}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".webm"}
ALLOWED_EXTENSIONS = IMAGE_EXTENSIONS | VIDEO_EXTENSIONS | {".pdf", ".pptx"}


def utc_now():
    return datetime.now(timezone.utc).isoformat()


def sha256(data):
    return hashlib.sha256(data).hexdigest()


def find_details(text):
    patterns = {
        "Email addresses": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        "Booking / PNR numbers": r"\b(?:PNR|booking(?:\s*(?:id|no|number|reference))?)\s*[:#-]?\s*[A-Z0-9-]{4,}\b",
        "Ticket numbers": r"\b(?:ticket(?:\s*(?:id|no|number))?)\s*[:#-]?\s*[A-Z0-9-]{4,}\b",
        "Transaction numbers": r"\b(?:transaction(?:\s*(?:id|no|number))?|UTR|reference\s*no)\s*[:#-]?\s*[A-Z0-9-]{4,}\b",
        "Dates": r"\b(?:\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{4}[/-]\d{1,2}[/-]\d{1,2})\b",
        "Times": r"\b(?:[01]?\d|2[0-3]):[0-5]\d(?:\s*[AP]M)?\b",
        "Amounts": r"(?:₹|INR|Rs\.?|USD|\$)\s?\d[\d,]*(?:\.\d{1,2})?",
        "Seat numbers": r"\bseat\s*(?:no\.?|number)?\s*[:#-]?\s*[A-Z0-9-]+\b",
    }
    found = {}
    for label, pattern in patterns.items():
        matches = list(dict.fromkeys(re.findall(pattern, text, re.IGNORECASE)))
        if matches:
            found[label] = matches[:30]
    return found


def read_ocr(image):
    """Return OCR text, or a clear message when Tesseract is not installed."""
    try:
        return pytesseract.image_to_string(image, config="--psm 6").strip(), None
    except pytesseract.pytesseract.TesseractNotFoundError:
        return "", (
            "OCR engine Tesseract was not found. Install Tesseract OCR for "
            "Windows, then restart this backend."
        )
    except Exception as exc:
        return "", f"OCR could not read this image: {str(exc)[:250]}"


def decode_qr(image):
    detector = cv2.QRCodeDetector()
    codes = []
    try:
        ok, texts, points, _ = detector.detectAndDecodeMulti(image)
        if ok and points is not None:
            for i, value in enumerate(texts):
                codes.append({
                    "text": value or "",
                    "readable": bool(value),
                    "corners": points[i].astype(float).tolist()
                    if i < len(points) else []
                })
    except (cv2.error, AttributeError):
        pass

    if not codes:
        try:
            value, points, _ = detector.detectAndDecode(image)
            if points is not None:
                codes.append({
                    "text": value or "",
                    "readable": bool(value),
                    "corners": points.astype(float).tolist()
                })
        except cv2.error:
            pass

    return {
        "status": "decoded" if any(c["readable"] for c in codes)
                  else "detected but unreadable" if codes
                  else "not detected",
        "count": len(codes),
        "codes": codes
    }


def image_analysis(data):
    with Image.open(io.BytesIO(data)) as original:
        original.verify()

    with Image.open(io.BytesIO(data)) as original:
        img = original.convert("RGB")
        width, height = img.size
        image_format = original.format or "Unknown"
        exif = {}
        try:
            for tag, value in original.getexif().items():
                name = ExifTags.TAGS.get(tag, str(tag))
                exif[name] = str(value)[:300]
        except Exception:
            pass

    rgb = np.array(img)
    bgr = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)
    ocr_text, ocr_error = read_ocr(img)
    qr = decode_qr(bgr)
    gray = cv2.cvtColor(rgb, cv2.COLOR_RGB2GRAY)
    sharpness = round(float(cv2.Laplacian(gray, cv2.CV_64F).var()), 2)

    ela_image = None
    signals = [{
        "name": "Image dimensions",
        "status": "informational",
        "detail": f"{width} × {height} pixels; format: {image_format}."
    }, {
        "name": "Image sharpness",
        "status": "review" if sharpness < 20 else "informational",
        "detail": f"Sharpness measurement: {sharpness}. Blur can affect OCR."
    }, {
        "name": "QR analysis",
        "status": qr["status"],
        "detail": f"{qr['count']} QR code region(s) detected."
    }]

    if ocr_error:
        signals.append({"name": "OCR", "status": "unavailable", "detail": ocr_error})
    else:
        signals.append({
            "name": "OCR text extraction",
            "status": "available" if ocr_text else "no text found",
            "detail": f"{len(ocr_text)} characters extracted from the image."
        })

    if image_format.upper() in ("JPEG", "JPG"):
        try:
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=90)
            buf.seek(0)
            with Image.open(buf) as compressed:
                other = np.array(compressed.convert("RGB"))
            difference = cv2.absdiff(rgb, other)
            heat = cv2.normalize(
                cv2.cvtColor(difference, cv2.COLOR_RGB2GRAY),
                None, 0, 255, cv2.NORM_MINMAX
            )
            _, encoded = cv2.imencode(".jpg", cv2.applyColorMap(heat, cv2.COLORMAP_JET))
            ela_image = base64.b64encode(encoded.tobytes()).decode("ascii")
            signals.append({
                "name": "Error Level Analysis",
                "status": "informational",
                "detail": "A recompression heatmap was generated. It is not proof of editing."
            })
        except Exception as exc:
            signals.append({"name": "Error Level Analysis", "status": "unavailable",
                            "detail": str(exc)[:200]})

    return {
        "file_type": "image",
        "image_details": {
            "format": image_format, "width": width, "height": height,
            "sharpness": sharpness
        },
        "exif_metadata": exif,
        "ocr_text": ocr_text,
        "detected_details": find_details(ocr_text),
        "ocr_error": ocr_error,
        "qr_analysis": qr,
        "ela_image": ela_image,
        "signals": signals,
        "limitations": [
            "OCR can misread blurred, rotated or low-quality text.",
            "ELA and image signals are clues, not proof of manipulation.",
            "A readable QR code does not prove a ticket or payment is genuine.",
            "Confirm authenticity through the official issuer or provider."
        ]
    }


def pdf_analysis(data):
    pages = []
    all_text = []
    ocr_errors = []

    with fitz.open(stream=data, filetype="pdf") as pdf:
        metadata = {}
        for key, value in (pdf.metadata or {}).items():
            if value not in (None, ""):
                metadata[key] = str(value)[:1000]

        page_count = len(pdf)
        pages_to_process = min(page_count, 30)

        for index in range(pages_to_process):
            page = pdf[index]
            text = (page.get_text("text") or "").strip()
            extraction_method = "embedded PDF text"

            # OCR scanned pages when normal PDF text is absent or extremely short.
            ocr_error = None
            if len(text) < 25:
                try:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    page_image = Image.open(io.BytesIO(pix.tobytes("png"))).convert("RGB")
                    scanned_text, ocr_error = read_ocr(page_image)
                    if len(scanned_text) > len(text):
                        text = scanned_text
                        extraction_method = "OCR"
                except Exception as exc:
                    ocr_error = f"Could not render this PDF page: {str(exc)[:200]}"

            if ocr_error:
                ocr_errors.append(f"Page {index + 1}: {ocr_error}")

            all_text.append(text)
            pages.append({
                "page": index + 1,
                "characters_extracted": len(text),
                "method": extraction_method,
                "text": text[:12000]
            })

        full_text = "\n\n".join(all_text)
        author_fields = {
            key: metadata[key]
            for key in ("author", "creator", "producer", "creationDate",
                        "modDate", "title", "subject", "keywords", "format",
                        "encryption") if metadata.get(key)
        }

    signals = [{
        "name": "PDF pages",
        "status": "informational",
        "detail": f"{pages_to_process} of {page_count} page(s) inspected."
    }, {
        "name": "Text extraction",
        "status": "available" if full_text.strip() else "no text found",
        "detail": f"{len(full_text)} characters extracted from the inspected pages."
    }, {
        "name": "PDF metadata",
        "status": "available" if metadata else "not found",
        "detail": f"{len(metadata)} metadata field(s) found."
    }]

    if ocr_errors:
        signals.append({
            "name": "OCR warnings", "status": "review",
            "detail": "Some pages could not be OCR-processed. See the page results and warnings."
        })

    return {
        "file_type": "pdf",
        "pdf_details": {
            "pages": page_count,
            "pages_inspected": pages_to_process,
            "metadata": metadata,
            "author_creator_details": author_fields,
            "page_summaries": pages,
            "ocr_warnings": ocr_errors
        },
        "ocr_text": full_text[:100000],
        "detected_details": find_details(full_text),
        "signals": signals,
        "limitations": [
            "PDF metadata can be missing or changed and does not prove who created a document.",
            "At most 30 pages are processed in this prototype.",
            "OCR may make mistakes; compare important details with the original.",
            "Ticket authenticity requires confirmation from the official issuer."
        ]
    }


def pptx_analysis(data):
    from pptx import Presentation
    presentation = Presentation(io.BytesIO(data))
    slides = []
    all_text = []
    for index, slide in enumerate(presentation.slides):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                parts.append(shape.text.strip())
        slide_text = "\n".join(parts)
        all_text.append(slide_text)
        slides.append({"slide": index + 1, "text": slide_text[:10000]})
    full_text = "\n\n".join(all_text)
    return {
        "file_type": "pptx",
        "pptx_details": {"slides": len(slides), "slide_summaries": slides},
        "ocr_text": full_text[:100000],
        "detected_details": find_details(full_text),
        "signals": [{"name": "PowerPoint text",
                     "status": "available" if full_text else "no text found",
                     "detail": f"{len(slides)} slides inspected."}],
        "limitations": ["Text inside slide images is not OCR-processed in this version."]
    }


def video_analysis(data, extension):
    import os
    import tempfile
    path = None
    video = None
    try:
        with tempfile.NamedTemporaryFile(suffix=extension, delete=False) as f:
            f.write(data)
            path = f.name
        video = cv2.VideoCapture(path)
        if not video.isOpened():
            raise ValueError("Could not open this video. Try MP4, MOV, AVI, MKV or WEBM.")
        fps = video.get(cv2.CAP_PROP_FPS) or 0
        frames = int(video.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
        duration = frames / fps if fps > 0 else None
        indices = np.linspace(0, max(0, frames - 1), min(10, frames), dtype=int) if frames else []
        samples, texts = [], []
        for index in indices:
            video.set(cv2.CAP_PROP_POS_FRAMES, int(index))
            ok, frame = video.read()
            if not ok:
                continue
            rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            text, _ = read_ocr(Image.fromarray(rgb))
            qr = decode_qr(frame)
            samples.append({
                "frame": int(index),
                "timestamp_seconds": round(index / fps, 2) if fps else None,
                "ocr_text": text[:1500],
                "qr_codes": qr["codes"]
            })
            if text:
                texts.append(text)
        full_text = "\n\n".join(texts)
        return {
            "file_type": "video",
            "video_details": {
                "fps": round(fps, 2), "frame_count": frames,
                "duration_seconds": round(duration, 2) if duration is not None else None,
                "sampled_frames": samples
            },
            "ocr_text": full_text[:20000],
            "detected_details": find_details(full_text),
            "signals": [{"name": "Video sampling", "status": "informational",
                         "detail": f"{len(samples)} frame(s) sampled for text and QR codes."}],
            "limitations": [
                "Only selected frames are inspected; edits between frames may be missed.",
                "Frame sampling does not prove video tampering."
            ]
        }
    finally:
        if video is not None:
            video.release()
        if path and Path(path).exists():
            Path(path).unlink(missing_ok=True)


def save_record(record):
    previous_hash = None
    if LOG_PATH.exists():
        with LOG_PATH.open("r", encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    previous_hash = json.loads(line).get("record_hash")

    record["previous_record_hash"] = previous_hash
    raw = json.dumps(record, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    record["record_hash"] = sha256(raw.encode("utf-8"))
    with LOG_PATH.open("a", encoding="utf-8") as f:
        f.write(json.dumps(record, ensure_ascii=False) + "\n")
    return record


@app.get("/")
def home():
    return jsonify({"name": "EvidenceGuard AI", "status": "running"})


@app.get("/api/health")
def health():
    try:
        pytesseract.get_tesseract_version()
        ocr_ready = True
        ocr_message = "Tesseract OCR is ready."
    except Exception:
        ocr_ready = False
        ocr_message = "Tesseract OCR is not installed or its path is not configured."
    return jsonify({
        "status": "running",
        "ocr_ready": ocr_ready,
        "ocr_message": ocr_message
    })


@app.post("/api/analyze")
def analyze():
    if "file" not in request.files:
        return jsonify({"error": "No file was received. Please choose a file."}), 400

    uploaded = request.files["file"]
    filename = Path(uploaded.filename or "captured-image.jpg").name
    extension = Path(filename).suffix.lower()
    data = uploaded.read()

    if not data:
        return jsonify({"error": "The selected file is empty."}), 400
    if extension not in ALLOWED_EXTENSIONS:
        return jsonify({"error": f"Unsupported file type: {extension or 'unknown'}"}), 400

    try:
        if extension in IMAGE_EXTENSIONS:
            result = image_analysis(data)
        elif extension == ".pdf":
            result = pdf_analysis(data)
        elif extension == ".pptx":
            result = pptx_analysis(data)
        else:
            result = video_analysis(data, extension)

        timestamp = utc_now()
        record = save_record({
            "evidence_id": "EV-" + uuid.uuid4().hex[:12].upper(),
            "action": "FILE_ANALYZED",
            "timestamp_utc": timestamp,
            "filename": filename,
            "file_type": result["file_type"],
            "file_size_bytes": len(data),
            "sha256": sha256(data)
        })

        result.update({
            "filename": filename,
            "file_size_bytes": len(data),
            "sha256": sha256(data),
            "analysis_timestamp_utc": timestamp,
            "evidence_record": record,
            "summary": (
                "Analysis completed. Review the extracted text and technical signals. "
                "These results alone cannot determine authenticity."
            )
        })
        return jsonify(result)

    except Exception as exc:
        app.logger.exception("File analysis failed")
        return jsonify({
            "error": f"Could not analyze this file: {str(exc)[:500]}"
        }), 500


@app.get("/api/history")
def history():
    records = []
    if LOG_PATH.exists():
        with LOG_PATH.open("r", encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    records.append(json.loads(line))
    return jsonify({"records": records[::-1][:100]})


@app.post("/api/verify-log")
def verify_log():
    records = []
    errors = []
    if LOG_PATH.exists():
        with LOG_PATH.open("r", encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    records.append(json.loads(line))

    previous = None
    for index, record in enumerate(records, start=1):
        saved_hash = record.get("record_hash")
        body = {k: v for k, v in record.items() if k != "record_hash"}
        raw = json.dumps(body, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
        if sha256(raw.encode("utf-8")) != saved_hash:
            errors.append(f"Record {index}: record hash mismatch.")
        if body.get("previous_record_hash") != previous:
            errors.append(f"Record {index}: previous-record link mismatch.")
        previous = saved_hash

    return jsonify({
        "valid": not errors,
        "records_checked": len(records),
        "errors": errors
    })


@app.errorhandler(413)
def too_large(_error):
    return jsonify({"error": "File is too large. Maximum allowed size is 100 MB."}), 413


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=True)
